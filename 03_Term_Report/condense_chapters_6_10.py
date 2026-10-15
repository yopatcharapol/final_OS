#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
สคริปต์ย่อเนื้อหาบทที่ 6-10 ของรายงาน Operating Systems
โดยควบคุมจำนวนหน้ารวมของบทที่ 6 ถึง 10 ให้อยู่ในช่วง 25 - 30 หน้าพอดีตามข้อกำหนด
รักษาความสมบูรณ์ทางวิชาการ ตารางผลการทดลองหลักทั้งหมด และรูปแบบฟอนต์ TH Sarabun New 16pt มาตรฐาน
"""

import os
import re
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

try:
    from pythainlp.tokenize import word_tokenize
    HAS_PYTHAINLP = True
except ImportError:
    HAS_PYTHAINLP = False

FONT_FAMILY = "TH Sarabun New"
CODE_FONT_FAMILY = "Consolas"

def insert_thai_zwsp(text):
    if not HAS_PYTHAINLP or not text or not any('\u0e00' <= char <= '\u0e7f' for char in text):
        return text
    words = word_tokenize(text, engine='newmm')
    result = []
    for i, w in enumerate(words):
        result.append(w)
        if i < len(words) - 1:
            next_w = words[i + 1]
            if not w.isspace() and not next_w.isspace() and not w.endswith('\n') and not next_w.startswith('\n'):
                result.append('\u200b')
    return "".join(result)

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_shading(cell, color_hex="F2F2F2"):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_table_borders(table, color="B8B8B8", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:insideV w:val="none"/>\n'
        f'  <w:left w:val="none"/>\n'
        f'  <w:right w:val="none"/>\n'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def apply_font_to_run(run, font_name=FONT_FAMILY, size_pt=16, bold=False, italic=False, color_rgb=None):
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    if color_rgb:
        run.font.color.rgb = color_rgb
    
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:ascii'), font_name)
    rFonts.set(qn('w:hAnsi'), font_name)
    rFonts.set(qn('w:cs'), font_name)
    rFonts.set(qn('w:eastAsia'), font_name)
    
    sz_half_pts = str(int(size_pt * 2))
    sz = rPr.find(qn('w:sz'))
    if sz is None:
        sz = OxmlElement('w:sz')
        rPr.append(sz)
    sz.set(qn('w:val'), sz_half_pts)
    
    szCs = rPr.find(qn('w:szCs'))
    if szCs is None:
        szCs = OxmlElement('w:szCs')
        rPr.append(szCs)
    szCs.set(qn('w:val'), sz_half_pts)
    
    if bold and rPr.find(qn('w:bCs')) is None:
        rPr.append(OxmlElement('w:bCs'))
    if italic and rPr.find(qn('w:iCs')) is None:
        rPr.append(OxmlElement('w:iCs'))
        
    lang = rPr.find(qn('w:lang'))
    if lang is None:
        lang = OxmlElement('w:lang')
        rPr.append(lang)
    lang.set(qn('w:val'), 'th-TH')
    lang.set(qn('w:bidi'), 'th-TH')
    lang.set(qn('w:eastAsia'), 'th-TH')

def format_paragraph(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, line_spacing=1.15,
                     space_before_pt=0, space_after_pt=4, first_line_indent_in=0.5,
                     keep_with_next=False):
    p.alignment = align
    p.paragraph_format.line_spacing = line_spacing
    p.paragraph_format.space_before = Pt(space_before_pt)
    p.paragraph_format.space_after = Pt(space_after_pt)
    p.paragraph_format.widow_control = True
    p.paragraph_format.first_line_indent = Inches(first_line_indent_in) if first_line_indent_in > 0 else Inches(0)
    p.paragraph_format.keep_with_next = keep_with_next

def parse_inline_markdown(p, text, default_size=16):
    pattern = re.compile(r'(\*\*.*?\*\*|\*.*?\*|`.*?`|\$.*?\$|\[.*?\])')
    tokens = pattern.split(text)
    for token in tokens:
        if not token:
            continue
        if token.startswith('**') and token.endswith('**') and len(token) >= 4:
            sub = token[2:-2]
            r = p.add_run(insert_thai_zwsp(sub))
            apply_font_to_run(r, size_pt=default_size, bold=True)
        elif token.startswith('*') and token.endswith('*') and len(token) >= 2:
            sub = token[1:-1]
            r = p.add_run(insert_thai_zwsp(sub))
            apply_font_to_run(r, size_pt=default_size, italic=True)
        elif token.startswith('`') and token.endswith('`') and len(token) >= 2:
            sub = token[1:-1]
            r = p.add_run(sub)
            apply_font_to_run(r, font_name=CODE_FONT_FAMILY, size_pt=default_size - 1.5, bold=False, color_rgb=RGBColor(30, 30, 30))
        elif token.startswith('$') and token.endswith('$') and len(token) >= 2:
            sub = token[1:-1]
            r = p.add_run(sub)
            apply_font_to_run(r, size_pt=default_size, italic=True)
        else:
            r = p.add_run(insert_thai_zwsp(token))
            apply_font_to_run(r, size_pt=default_size, bold=False, italic=False)

def add_styled_paragraph(doc, text, style_type="body"):
    p = doc.add_paragraph()
    if style_type == "chapter":
        format_paragraph(p, align=WD_ALIGN_PARAGRAPH.CENTER, line_spacing=1.15,
                         space_before_pt=14, space_after_pt=14, first_line_indent_in=0,
                         keep_with_next=True)
        r = p.add_run(insert_thai_zwsp(text))
        apply_font_to_run(r, size_pt=20, bold=True)
    elif style_type == "heading1":
        format_paragraph(p, align=WD_ALIGN_PARAGRAPH.LEFT, line_spacing=1.15,
                         space_before_pt=10, space_after_pt=4, first_line_indent_in=0,
                         keep_with_next=True)
        r = p.add_run(insert_thai_zwsp(text))
        apply_font_to_run(r, size_pt=18, bold=True)
    elif style_type == "heading2":
        format_paragraph(p, align=WD_ALIGN_PARAGRAPH.LEFT, line_spacing=1.15,
                         space_before_pt=6, space_after_pt=4, first_line_indent_in=0,
                         keep_with_next=True)
        r = p.add_run(insert_thai_zwsp(text))
        apply_font_to_run(r, size_pt=16, bold=True)
    elif style_type == "table_caption":
        format_paragraph(p, align=WD_ALIGN_PARAGRAPH.LEFT, line_spacing=1.15,
                         space_before_pt=8, space_after_pt=4, first_line_indent_in=0,
                         keep_with_next=True)
        r = p.add_run(insert_thai_zwsp(text))
        apply_font_to_run(r, size_pt=16, bold=True)
    elif style_type == "figure_caption":
        format_paragraph(p, align=WD_ALIGN_PARAGRAPH.CENTER, line_spacing=1.15,
                         space_before_pt=4, space_after_pt=8, first_line_indent_in=0,
                         keep_with_next=False)
        r = p.add_run(insert_thai_zwsp(text))
        apply_font_to_run(r, size_pt=16, bold=True)
    elif style_type == "bullet":
        format_paragraph(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, line_spacing=1.15,
                         space_before_pt=0, space_after_pt=3, first_line_indent_in=0.25,
                         keep_with_next=False)
        parse_inline_markdown(p, text, default_size=16)
    else: # body
        format_paragraph(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, line_spacing=1.15,
                         space_before_pt=0, space_after_pt=4, first_line_indent_in=0.5,
                         keep_with_next=False)
        parse_inline_markdown(p, text, default_size=16)
    return p

def add_code_block(doc, code_lines):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_shading(cell, "F8F9FA")
    set_cell_margins(cell, top=120, bottom=120, left=160, right=160)
    
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="single" w:sz="4" w:space="0" w:color="D0D5DD"/>\n'
        f'  <w:left w:val="single" w:sz="12" w:space="0" w:color="0052CC"/>\n'
        f'  <w:bottom w:val="single" w:sz="4" w:space="0" w:color="D0D5DD"/>\n'
        f'  <w:right w:val="single" w:sz="4" w:space="0" w:color="D0D5DD"/>\n'
        f'</w:tcBorders>'
    )
    tcPr.append(tcBorders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.first_line_indent = Inches(0)
    
    code_text = "\n".join(code_lines)
    r = p.add_run(code_text)
    apply_font_to_run(r, font_name=CODE_FONT_FAMILY, size_pt=9.5, bold=False, color_rgb=RGBColor(33, 37, 41))

    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.space_after = Pt(4)
    sp.paragraph_format.line_spacing = 1.0

def add_markdown_table(doc, table_rows, col_widths=None):
    if not table_rows or len(table_rows) < 2:
        return
    
    header_cols = [c.strip() for c in table_rows[0].split('|')[1:-1]]
    data_rows = []
    for r in table_rows[1:]:
        if re.match(r'^\s*\|?\s*[-:]+[-| :]*\|?\s*$', r):
            continue
        cols = [c.strip() for c in r.split('|')[1:-1]]
        if cols:
            data_rows.append(cols)
            
    num_cols = len(header_cols)
    if num_cols == 0:
        return
        
    tbl = doc.add_table(rows=len(data_rows) + 1, cols=num_cols)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl, color="B8B8B8", sz="4", val="single")
    
    total_width_in = 5.77
    if not col_widths or len(col_widths) != num_cols:
        col_widths = [total_width_in / num_cols] * num_cols
        
    hdr_cells = tbl.rows[0].cells
    tbl.rows[0]._tr.get_or_add_trPr().append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
    tbl.rows[0]._tr.get_or_add_trPr().append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
    
    for i, title in enumerate(header_cols):
        if i < len(hdr_cells):
            cell = hdr_cells[i]
            cell.width = Inches(col_widths[i])
            set_cell_shading(cell, "EAECEE")
            set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
            p = cell.paragraphs[0]
            format_paragraph(p, align=WD_ALIGN_PARAGRAPH.CENTER, line_spacing=1.0,
                             space_before_pt=2, space_after_pt=2, first_line_indent_in=0)
            parse_inline_markdown(p, title, default_size=14)
            for run in p.runs:
                run.bold = True
                
    for r_idx, row_data in enumerate(data_rows):
        row = tbl.rows[r_idx + 1]
        row._tr.get_or_add_trPr().append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        shading = "FFFFFF" if r_idx % 2 == 0 else "F8F9F9"
        
        for c_idx, cell_value in enumerate(row_data):
            if c_idx < len(row.cells):
                cell = row.cells[c_idx]
                cell.width = Inches(col_widths[c_idx])
                set_cell_shading(cell, shading)
                set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
                p = cell.paragraphs[0]
                align = WD_ALIGN_PARAGRAPH.CENTER if (len(cell_value) < 14 and any(char.isdigit() for char in cell_value)) else WD_ALIGN_PARAGRAPH.LEFT
                format_paragraph(p, align=align, line_spacing=1.05,
                                 space_before_pt=2, space_after_pt=2, first_line_indent_in=0)
                parse_inline_markdown(p, cell_value, default_size=13.5)

    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(2)
    sp.paragraph_format.space_after = Pt(4)
    sp.paragraph_format.line_spacing = 1.0

def setup_document_defaults(doc):
    for section in doc.sections:
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.top_margin = Inches(1.5)
        section.left_margin = Inches(1.5)
        section.bottom_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
    normal_style = doc.styles['Normal']
    normal_style.font.name = FONT_FAMILY
    normal_style.font.size = Pt(16)

print("Helper functions defined.")

