#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
สคริปต์สร้างเอกสาร Word (.docx) มาตรฐานวิชาการสำหรับรายงาน
เรื่อง: Thread & Multithreading ในระบบปฏิบัติการสมัยใหม่
(Thread Management and Its Performance Implications in Modern Operating Systems:
 Architecture, Synchronization, and Empirical Evaluation)

การตั้งค่าหน้ากระดาษและฟอนต์ตามมาตรฐาน (ฉบับแก้ไขการตัดคำและช่องไฟภาษาไทยสมบูรณ์):
- ฟอนต์: TH Sarabun New ทั้งเอกสาร (Complex Script + Latin + EastAsia)
- ใช้ PyThaiNLP แทรก Zero-Width Space (U+200B) เพื่อให้ Word ตัดคำภาษาไทยตามพยางค์และคำอย่างสมบูรณ์
- จัดขอบชิดสองด้าน (Justified: w:val="both") ปราศจากช่องว่างกว้างผิดปกติระหว่างคำ
- กำหนดภาษาของเอกสารเป็น th-TH และค่าขนาดตัวอักษร Complex Script (szCs) ให้เท่ากับ 16 pt
- ระยะขอบกระดาษ: บน 1.5 นิ้ว, ซ้าย 1.5 นิ้ว (สำหรับเข้าเล่ม), ล่าง 1.0 นิ้ว, ขวา 1.0 นิ้ว
- ชื่อบท: ขนาด 20 pt หนา จัดกึ่งกลางหน้า
- หัวข้อหลัก: ขนาด 18 pt หนา ชิดซ้าย
- หัวข้อย่อย: ขนาด 16 pt หนา ชิดซ้าย
- เนื้อความทั่วไป: ขนาด 16 pt ปกติ ย่อหน้าบรรทัดแรก 0.5 นิ้ว (1.27 ซม.) จัดชิดขอบสองด้าน (Justified)
- ระยะห่างบรรทัด: 1.15 บรรทัด และเว้นระยะหลังย่อหน้า 4 pt
- ชื่อตาราง: ขนาด 16 pt หนา ด้านบนตาราง ชิดซ้าย
- ชื่อรูปภาพ: ขนาด 16 pt หนา ด้านล่างรูปภาพ กึ่งกลาง
- ส่งออกเป็นไฟล์: Report_Final.docx (พร้อมไฟล์สำรอง Report_Final_Fixed.docx)
"""

import os
import re
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

try:
    from pythainlp.tokenize import word_tokenize
    HAS_PYTHAINLP = True
except ImportError:
    HAS_PYTHAINLP = False

FONT_FAMILY = "TH Sarabun New"
CODE_FONT_FAMILY = "Consolas"

def insert_thai_zwsp(text):
    """
    แทรก Zero-Width Space (U+200B) ที่รอยต่อคำภาษาไทย
    เพื่อให้ Microsoft Word สามารถตัดคำขึ้นบรรทัดใหม่ได้อย่างแม่นยำ ป้องกันการเกิดช่องว่างขนาดใหญ่
    """
    if not HAS_PYTHAINLP or not text or not any('\u0e00' <= char <= '\u0e7f' for char in text):
        return text
    words = word_tokenize(text, engine='newmm')
    result = []
    for i, w in enumerate(words):
        result.append(w)
        if i < len(words) - 1:
            next_w = words[i + 1]
            # แทรก \u200b เฉพาะเมื่อระหว่าง w และ next_w ไม่ใช่ space หรือ newline
            if not w.isspace() and not next_w.isspace() and not w.endswith('\n') and not next_w.startswith('\n'):
                result.append('\u200b')
    return "".join(result)

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    """กำหนดระยะขอบภายในเซลล์ตาราง (หน่วย twips: 20 twips = 1 pt)"""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_shading(cell, color_hex="F2F2F2"):
    """กำหนดสีพื้นหลังของเซลล์ตาราง"""
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_table_borders(table, color="B8B8B8", sz="4", val="single"):
    """กำหนดเส้นขอบตารางให้เรียบร้อยสวยงามแบบวิชาการ"""
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:insideV w:val="none"/>\n'
        f'  <w:left w:val="none"/>\n'
        f'  <w:right w:val="none"/>\n'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def apply_font_to_run(run, font_name=FONT_FAMILY, size_pt=16, bold=False, italic=False, color_rgb=None):
    """กำหนดฟอนต์ให้กับ Run ทั้งส่วน Latin และ Complex Script (พร้อมภาษา th-TH)"""
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    if color_rgb:
        run.font.color.rgb = color_rgb
    
    rPr = run._r.get_or_add_rPr()
    
    # 1. กำหนดชื่อฟอนต์ครอบคลุมทุกระบบสคริปต์
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:ascii'), font_name)
    rFonts.set(qn('w:hAnsi'), font_name)
    rFonts.set(qn('w:cs'), font_name)
    rFonts.set(qn('w:eastAsia'), font_name)
    
    # 2. กำหนดขนาดฟอนต์สำหรับ Complex Script (szCs) ให้เท่ากับขนาดปกติ (sz)
    sz_half_pts = str(int(size_pt * 2))
    sz = rPr.find(qn('w:sz'))
    if sz is None:
        sz = OxmlElement('w:sz')
        rPr.append(sz)
    sz.set(qn('w:val'), sz_half_pts)
    
    szCs = rPr.find(qn('w:szCs'))
    if szCs is None:
        szCs = OxmlElement('w:szCs')
        rPr.append(szCs)
    szCs.set(qn('w:val'), sz_half_pts)
    
    # 3. กำหนดตัวหนา/ตัวเอียงสำหรับ Complex Script
    if bold:
        if rPr.find(qn('w:bCs')) is None:
            rPr.append(OxmlElement('w:bCs'))
    if italic:
        if rPr.find(qn('w:iCs')) is None:
            rPr.append(OxmlElement('w:iCs'))
            
    # 4. กำหนดภาษาเป็นภาษาไทย th-TH เพื่อให้ Word ตัดคำอย่างเป็นธรรมชาติ
    lang = rPr.find(qn('w:lang'))
    if lang is None:
        lang = OxmlElement('w:lang')
        rPr.append(lang)
    lang.set(qn('w:val'), 'th-TH')
    lang.set(qn('w:bidi'), 'th-TH')
    lang.set(qn('w:eastAsia'), 'th-TH')

def format_paragraph(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, line_spacing=1.15,
                     space_before_pt=0, space_after_pt=4, first_line_indent_in=0.5,
                     keep_with_next=False):
    """
    ตั้งค่าพารากราฟตามเกณฑ์รายงานวิชาการ
    ใช้ WD_ALIGN_PARAGRAPH.JUSTIFY (w:val="both") เพื่อป้องกันการยืดช่องไฟตัวอักษรแบบ thaiDistribute
    """
    p.alignment = align
    p.paragraph_format.line_spacing = line_spacing
    p.paragraph_format.space_before = Pt(space_before_pt)
    p.paragraph_format.space_after = Pt(space_after_pt)
    p.paragraph_format.widow_control = True
    if first_line_indent_in > 0:
        p.paragraph_format.first_line_indent = Inches(first_line_indent_in)
    else:
        p.paragraph_format.first_line_indent = Inches(0)
    p.paragraph_format.keep_with_next = keep_with_next

def add_styled_paragraph(doc, text, style_type="body"):
    """สร้างย่อหน้าพร้อมกำหนดสไตล์ตามประเภทเนื้อหา"""
    p = doc.add_paragraph()
    
    if style_type == "chapter":
        format_paragraph(p, align=WD_ALIGN_PARAGRAPH.CENTER, line_spacing=1.15,
                         space_before_pt=14, space_after_pt=14, first_line_indent_in=0,
                         keep_with_next=True)
        r = p.add_run(insert_thai_zwsp(text))
        apply_font_to_run(r, size_pt=20, bold=True)
        return p

    elif style_type == "heading1":
        format_paragraph(p, align=WD_ALIGN_PARAGRAPH.LEFT, line_spacing=1.15,
                         space_before_pt=10, space_after_pt=4, first_line_indent_in=0,
                         keep_with_next=True)
        r = p.add_run(insert_thai_zwsp(text))
        apply_font_to_run(r, size_pt=18, bold=True)
        return p

    elif style_type == "heading2":
        format_paragraph(p, align=WD_ALIGN_PARAGRAPH.LEFT, line_spacing=1.15,
                         space_before_pt=6, space_after_pt=4, first_line_indent_in=0,
                         keep_with_next=True)
        r = p.add_run(insert_thai_zwsp(text))
        apply_font_to_run(r, size_pt=16, bold=True)
        return p

    elif style_type == "table_caption":
        format_paragraph(p, align=WD_ALIGN_PARAGRAPH.LEFT, line_spacing=1.15,
                         space_before_pt=8, space_after_pt=4, first_line_indent_in=0,
                         keep_with_next=True)
        r = p.add_run(insert_thai_zwsp(text))
        apply_font_to_run(r, size_pt=16, bold=True)
        return p

    elif style_type == "figure_caption":
        format_paragraph(p, align=WD_ALIGN_PARAGRAPH.CENTER, line_spacing=1.15,
                         space_before_pt=4, space_after_pt=8, first_line_indent_in=0,
                         keep_with_next=False)
        r = p.add_run(insert_thai_zwsp(text))
        apply_font_to_run(r, size_pt=16, bold=True)
        return p

    elif style_type == "equation":
        format_paragraph(p, align=WD_ALIGN_PARAGRAPH.CENTER, line_spacing=1.15,
                         space_before_pt=6, space_after_pt=6, first_line_indent_in=0,
                         keep_with_next=False)
        r = p.add_run(text)
        apply_font_to_run(r, size_pt=16, italic=True)
        return p

    elif style_type == "bullet":
        format_paragraph(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, line_spacing=1.15,
                         space_before_pt=0, space_after_pt=3, first_line_indent_in=0.25,
                         keep_with_next=False)
        parse_inline_markdown(p, text, default_size=16)
        return p

    else: # body
        format_paragraph(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, line_spacing=1.15,
                         space_before_pt=0, space_after_pt=4, first_line_indent_in=0.5,
                         keep_with_next=False)
        parse_inline_markdown(p, text, default_size=16)
        return p

def parse_inline_markdown(p, text, default_size=16):
    """ตัดแยกโทเค็นตัวหนา ตัวเอียง และโค้ดในข้อความแล้วใส่ Run ที่จัดฟอนต์ถูกต้อง"""
    pattern = re.compile(r'(\*\*.*?\*\*|\*.*?\*|`.*?`|\$.*?\$|\[.*?\])')
    tokens = pattern.split(text)
    
    for token in tokens:
        if not token:
            continue
        if token.startswith('**') and token.endswith('**') and len(token) >= 4:
            sub = token[2:-2]
            r = p.add_run(insert_thai_zwsp(sub))
            apply_font_to_run(r, size_pt=default_size, bold=True)
        elif token.startswith('*') and token.endswith('*') and len(token) >= 2:
            sub = token[1:-1]
            r = p.add_run(insert_thai_zwsp(sub))
            apply_font_to_run(r, size_pt=default_size, italic=True)
        elif token.startswith('`') and token.endswith('`') and len(token) >= 2:
            sub = token[1:-1]
            r = p.add_run(sub)
            apply_font_to_run(r, font_name=CODE_FONT_FAMILY, size_pt=default_size - 1.5, bold=False, color_rgb=RGBColor(30, 30, 30))
        elif token.startswith('$') and token.endswith('$') and len(token) >= 2:
            sub = token[1:-1]
            r = p.add_run(sub)
            apply_font_to_run(r, size_pt=default_size, italic=True)
        else:
            r = p.add_run(insert_thai_zwsp(token))
            apply_font_to_run(r, size_pt=default_size, bold=False, italic=False)

def add_code_block(doc, code_lines):
    """สร้างบล็อกโค้ดหรือไดอะแกรม ASCII ในกล่องที่จัดฟอนต์ Consolas เรียบร้อย"""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_shading(cell, "F8F9FA")
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
    
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="single" w:sz="4" w:space="0" w:color="D0D5DD"/>\n'
        f'  <w:left w:val="single" w:sz="12" w:space="0" w:color="0052CC"/>\n'
        f'  <w:bottom w:val="single" w:sz="4" w:space="0" w:color="D0D5DD"/>\n'
        f'  <w:right w:val="single" w:sz="4" w:space="0" w:color="D0D5DD"/>\n'
        f'</w:tcBorders>'
    )
    tcPr.append(tcBorders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.first_line_indent = Inches(0)
    
    code_text = "\n".join(code_lines)
    r = p.add_run(code_text)
    apply_font_to_run(r, font_name=CODE_FONT_FAMILY, size_pt=9.5, bold=False, color_rgb=RGBColor(33, 37, 41))

    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.space_after = Pt(4)
    sp.paragraph_format.line_spacing = 1.0

def add_markdown_table(doc, table_rows):
    """สร้างตาราง Word จากแถวตาราง Markdown พร้อมกำหนดความกว้างและตัดคำเรียบร้อย"""
    if not table_rows or len(table_rows) < 2:
        return
    
    header_cols = [c.strip() for c in table_rows[0].split('|')[1:-1]]
    data_rows = []
    for r in table_rows[1:]:
        if re.match(r'^\s*\|?\s*[-:]+[-| :]*\|?\s*$', r):
            continue
        cols = [c.strip() for c in r.split('|')[1:-1]]
        if cols:
            data_rows.append(cols)
            
    num_cols = len(header_cols)
    if num_cols == 0:
        return
        
    tbl = doc.add_table(rows=len(data_rows) + 1, cols=num_cols)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl, color="B8B8B8", sz="4", val="single")
    
    # กำหนดความกว้างรวมของตารางให้เท่ากับ 5.77 นิ้ว (ความกว้างหน้า A4 หักขอบซ้าย 1.5" และขวา 1.0")
    total_width_in = 5.77
    col_width = Inches(total_width_in / num_cols)
    
    # Format Header Row
    hdr_cells = tbl.rows[0].cells
    tbl.rows[0]._tr.get_or_add_trPr().append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
    tbl.rows[0]._tr.get_or_add_trPr().append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
    
    for i, title in enumerate(header_cols):
        if i < len(hdr_cells):
            cell = hdr_cells[i]
            cell.width = col_width
            set_cell_shading(cell, "EAECEE")
            set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
            p = cell.paragraphs[0]
            format_paragraph(p, align=WD_ALIGN_PARAGRAPH.CENTER, line_spacing=1.0,
                             space_before_pt=2, space_after_pt=2, first_line_indent_in=0)
            parse_inline_markdown(p, title, default_size=14.5)
            for run in p.runs:
                run.bold = True
                
    # Format Data Rows
    for r_idx, row_data in enumerate(data_rows):
        row = tbl.rows[r_idx + 1]
        row._tr.get_or_add_trPr().append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        shading = "FFFFFF" if r_idx % 2 == 0 else "F8F9F9"
        
        for c_idx, cell_value in enumerate(row_data):
            if c_idx < len(row.cells):
                cell = row.cells[c_idx]
                cell.width = col_width
                set_cell_shading(cell, shading)
                set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
                p = cell.paragraphs[0]
                align = WD_ALIGN_PARAGRAPH.CENTER if (len(cell_value) < 14 and any(char.isdigit() for char in cell_value)) else WD_ALIGN_PARAGRAPH.LEFT
                format_paragraph(p, align=align, line_spacing=1.1,
                                 space_before_pt=2, space_after_pt=2, first_line_indent_in=0)
                parse_inline_markdown(p, cell_value, default_size=14)

    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(2)
    sp.paragraph_format.space_after = Pt(4)
    sp.paragraph_format.line_spacing = 1.0

def create_title_page(doc):
    """สร้างหน้าปกรายงานเชิงวิชาการตามมาตรฐานสถาบัน"""
    p_top = doc.add_paragraph()
    format_paragraph(p_top, align=WD_ALIGN_PARAGRAPH.CENTER, line_spacing=1.15,
                     space_before_pt=36, space_after_pt=18, first_line_indent_in=0)
    r_type = p_top.add_run("รายงานเชิงวิชาการ\nวิชาระบบปฏิบัติการ (Operating Systems)")
    apply_font_to_run(r_type, size_pt=18, bold=True)
    
    p_title = doc.add_paragraph()
    format_paragraph(p_title, align=WD_ALIGN_PARAGRAPH.CENTER, line_spacing=1.2,
                     space_before_pt=40, space_after_pt=18, first_line_indent_in=0)
    r_th = p_title.add_run("เรื่อง\nThread & Multithreading ในระบบปฏิบัติการสมัยใหม่")
    apply_font_to_run(r_th, size_pt=22, bold=True)
    
    p_en = doc.add_paragraph()
    format_paragraph(p_en, align=WD_ALIGN_PARAGRAPH.CENTER, line_spacing=1.15,
                     space_before_pt=12, space_after_pt=48, first_line_indent_in=0)
    r_en = p_en.add_run("(Thread Management and Its Performance Implications in Modern Operating Systems:\nArchitecture, Synchronization, and Empirical Evaluation)")
    apply_font_to_run(r_en, size_pt=16, italic=True)
    
    p_bot = doc.add_paragraph()
    format_paragraph(p_bot, align=WD_ALIGN_PARAGRAPH.CENTER, line_spacing=1.2,
                     space_before_pt=120, space_after_pt=12, first_line_indent_in=0)
    r_bot = p_bot.add_run("สาขาวิชาวิศวกรรมซอฟต์แวร์ / วิทยาการคอมพิวเตอร์\nคณะเทคโนโลยีสารสนเทศและการสื่อสาร มหาวิทยาลัยพะเยา\nภาคเรียนที่ 1 ปีการศึกษา 2567")
    apply_font_to_run(r_bot, size_pt=16, bold=False)
    
    doc.add_page_break()

def is_ignore_line(line):
    s = line.strip()
    if not s:
        return True
    if s.startswith('# ต้นฉบับรายงานวิชา') or s.startswith('## ส่วนที่') or s.startswith('### หัวข้อ:'):
        return True
    if s.startswith('## 📌 คำแนะนำ') or s.startswith('1. **การจัดหน้า') or s.startswith('2. **การกะบรรทัด') or s.startswith('3. **การใช้สี') or s.startswith('4. **การย่อหน้า'):
        return True
    if s.startswith('* ขอบซ้าย') or s.startswith('* ขอบขวา') or s.startswith('* ขอบบน') or s.startswith('* กระดาษ A4') or s.startswith('* เนื้อหาใน') or s.startswith('* **สีน้ำเงิน') or s.startswith('* **สีดำ') or s.startswith('* **สีแดง') or s.startswith('* ทุกย่อหน้า'):
        return True
    if s.startswith('<!-- =================') or s.startswith('<!-- หน้าที่') or s.startswith('<!--'):
        return True
    if s.startswith('### [หน้าที่') or s.startswith('*(เริ่มเขียน') or s.startswith('*(ต่อเนื้อหา') or s.startswith('*(เริ่มหน้าใหม่') or s.startswith('*(เริ่มบทใหม่'):
        return True
    if s.startswith('*(สิ้นสุดเนื้อหา') or s.startswith('*(สิ้นสุดรายงาน') or s.startswith('*(ยอดรวม'):
        return True
    if s.startswith('*(เว้น 1 บรรทัด') or s.startswith('*(จัดทำตามมาตรฐาน') or s.startswith('*(พื้นที่วาดรูป') or s.startswith('*(ใช้ไม้บรรทัด'):
        return True
    if s == '---':
        return True
    return False

def clean_text_line(text):
    t = text.strip()
    t = re.sub(r'^\[ย่อหน้า\s*1\.5\s*ซม\.\]\s*', '', t)
    t = re.sub(r'^\[ย่อหน้า\s*1\s*ซม\.\]\s*', '', t)
    return t

def setup_document_defaults(doc):
    """ตั้งค่า Default Styles และภาษา th-TH ให้กับเอกสาร"""
    # 1. ตั้งค่าหน้ากระดาษ
    for section in doc.sections:
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.top_margin = Inches(1.5)
        section.left_margin = Inches(1.5)
        section.bottom_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
    # 2. ตั้งค่า Normal Style
    normal_style = doc.styles['Normal']
    normal_style.font.name = FONT_FAMILY
    normal_style.font.size = Pt(16)
    
    # 3. ตั้งค่า Default Languages and Fonts ใน docDefaults
    docDefaults = doc.styles.element.find(qn('w:docDefaults'))
    if docDefaults is not None:
        rPrDefault = docDefaults.find(qn('w:rPrDefault'))
        if rPrDefault is None:
            rPrDefault = OxmlElement('w:rPrDefault')
            docDefaults.append(rPrDefault)
        doc_rPr = rPrDefault.find(qn('w:rPr'))
        if doc_rPr is None:
            doc_rPr = OxmlElement('w:rPr')
            rPrDefault.append(doc_rPr)
            
        for child in list(doc_rPr):
            if child.tag.endswith('rFonts') or child.tag.endswith('lang'):
                doc_rPr.remove(child)
                
        rFonts = parse_xml(f'<w:rFonts {nsdecls("w")} w:ascii="{FONT_FAMILY}" w:hAnsi="{FONT_FAMILY}" w:cs="{FONT_FAMILY}" w:eastAsia="{FONT_FAMILY}"/>')
        lang = parse_xml(f'<w:lang {nsdecls("w")} w:val="th-TH" w:eastAsia="th-TH" w:bidi="th-TH"/>')
        sz = parse_xml(f'<w:sz {nsdecls("w")} w:val="32"/>')
        szCs = parse_xml(f'<w:szCs {nsdecls("w")} w:val="32"/>')
        doc_rPr.append(rFonts)
        doc_rPr.append(lang)
        doc_rPr.append(sz)
        doc_rPr.append(szCs)

def parse_manuscript_files(file_paths, output_docx_path):
    print(f"กำลังสร้างเอกสารฉบับแก้ไขการตัดคำ: {output_docx_path}")
    doc = Document()
    
    setup_document_defaults(doc)
    create_title_page(doc)
    
    chapter_count = 0
    in_code_block = False
    code_block_lines = []
    in_table = False
    table_lines = []
    
    for file_path in file_paths:
        if not os.path.exists(file_path):
            continue
            
        print(f"กำลังประมวลผลไฟล์: {os.path.basename(file_path)}")
        with open(file_path, 'r', encoding='utf-8') as f:
            lines = f.readlines()
            
        for line in lines:
            raw_line = line.rstrip('\r\n')
            stripped = raw_line.strip()
            
            if stripped.startswith('```'):
                if in_code_block:
                    add_code_block(doc, code_block_lines)
                    code_block_lines = []
                    in_code_block = False
                else:
                    if in_table:
                        add_markdown_table(doc, table_lines)
                        table_lines = []
                        in_table = False
                    in_code_block = True
                    code_block_lines = []
                continue
                
            if in_code_block:
                code_block_lines.append(raw_line)
                continue
                
            if stripped.startswith('|') and stripped.endswith('|'):
                in_table = True
                table_lines.append(stripped)
                continue
            else:
                if in_table:
                    add_markdown_table(doc, table_lines)
                    table_lines = []
                    in_table = False
                    
            if is_ignore_line(stripped):
                continue
                
            cleaned = clean_text_line(stripped)
            if not cleaned:
                continue
                
            if cleaned == '**Thread & Multithreading ในระบบปฏิบัติการสมัยใหม่**':
                continue
                
            # --- 1. ชื่อบท (Chapter Title) ---
            if (re.match(r'^\*\*บทที่\s*\d+:', cleaned) or 
                'เอกสารอ้างอิงและบรรณานุกรม' in cleaned or 
                re.match(r'^\*\*ภาคผนวก\s*[ก-ฮA-Z]:', cleaned)):
                clean_title = re.sub(r'^\*\*(.*?)\*\*$', r'\1', cleaned)
                chapter_count += 1
                if chapter_count > 1:
                    doc.add_page_break()
                add_styled_paragraph(doc, clean_title, style_type="chapter")
                continue
                
            # --- 2. หัวข้อหลัก (Main Heading 1.1, 1.2, 2.1 ...) ---
            m_h1 = re.match(r'^\*\*(\d+\.\d+)\s+(.*?)\*\*$', cleaned)
            m_ref_h1 = re.match(r'^\*\*(\d+)\.\s+(.*?)\*\*$', cleaned)
            if m_h1:
                clean_h1 = f"{m_h1.group(1)} {m_h1.group(2)}"
                add_styled_paragraph(doc, clean_h1, style_type="heading1")
                continue
            elif m_ref_h1 and ('ตำราเรียน' in cleaned or 'เอกสารมาตรฐาน' in cleaned or 'สถาปัตยกรรม' in cleaned or 'รายงานเชิง' in cleaned):
                clean_h1 = f"{m_ref_h1.group(1)}. {m_ref_h1.group(2)}"
                add_styled_paragraph(doc, clean_h1, style_type="heading1")
                continue
                
            # --- 3. หัวข้อย่อย (Subheading 1.1.1, 1.1.2 ...) ---
            m_h2 = re.match(r'^\*\*(\d+\.\d+\.\d+)\s+(.*?)\*\*$', cleaned)
            if m_h2:
                clean_h2 = f"{m_h2.group(1)} {m_h2.group(2)}"
                add_styled_paragraph(doc, clean_h2, style_type="heading2")
                continue
                
            # หัวข้อย่อยลำดับเลข เช่น **1. ภาวะติดตาย (Deadlock ในระดับเธรด)** หรือ **2. สปินล็อก...**
            m_h2_num = re.match(r'^\*\*(\d+)\.\s+(.*?)\*\*$', cleaned)
            if m_h2_num and len(cleaned) < 140:
                clean_h2 = f"{m_h2_num.group(1)}. {m_h2_num.group(2)}"
                add_styled_paragraph(doc, clean_h2, style_type="heading2")
                continue

            m_h1_other = re.match(r'^\*\*(\d+\.\d+)\s+(.*?)\b', cleaned)
            if m_h1_other and cleaned.startswith('**') and cleaned.endswith('**'):
                clean_h1 = cleaned[2:-2]
                add_styled_paragraph(doc, clean_h1, style_type="heading1")
                continue

            # หัวข้อย่อยข้อความตัวหนาแบบบรรทัดเดี่ยว เช่น **แนวทางการป้องกัน Deadlock...** หรือ **ปัญหา Livelock...**
            if cleaned.startswith('**') and cleaned.endswith('**') and len(cleaned) < 140 and not cleaned.startswith('**ตารางที่'):
                clean_subhead = cleaned[2:-2].strip()
                add_styled_paragraph(doc, clean_subhead, style_type="heading2")
                continue

            # --- 4. ชื่อตาราง (Table Caption) ---
            if cleaned.startswith('**ตารางที่') or cleaned.startswith('ตารางที่'):
                clean_caption = re.sub(r'^\*\*(.*?)\*\*$', r'\1', cleaned)
                add_styled_paragraph(doc, clean_caption, style_type="table_caption")
                continue

            # --- 5. ชื่อรูปภาพ (Figure Caption) ---
            if re.match(r'^รูปที่\s*\d+\.\d+', cleaned):
                add_styled_paragraph(doc, cleaned, style_type="figure_caption")
                continue

            # --- 6. สมการคณิตศาสตร์ (Standalone Equations) ---
            if cleaned.startswith('$$') and cleaned.endswith('$$'):
                eq_text = cleaned[2:-2].strip()
                add_styled_paragraph(doc, eq_text, style_type="equation")
                continue

            # --- 7. Bullet or Numbered List Items ---
            if re.match(r'^(\*|-|\d+\.)\s+', cleaned):
                add_styled_paragraph(doc, cleaned, style_type="bullet")
                continue

            # --- 8. Regular Body Paragraph ---
            add_styled_paragraph(doc, cleaned, style_type="body")

    if in_code_block and code_block_lines:
        add_code_block(doc, code_block_lines)
    if in_table and table_lines:
        add_markdown_table(doc, table_lines)

    try:
        doc.save(output_docx_path)
        print(f"สร้างเอกสารสำเร็จสมบูรณ์: {output_docx_path}")
    except PermissionError:
        alt_path = os.path.join(os.path.dirname(output_docx_path), 'Report_Final_Fixed.docx')
        doc.save(alt_path)
        print(f"ไฟล์ {os.path.basename(output_docx_path)} ถูกเปิดค้างอยู่ในโปรแกรม Microsoft Word")
        print(f"ระบบได้บันทึกไฟล์ฉบับแก้ไขสมบูรณ์ไว้ที่: {alt_path}")
        print("กรุณาปิดไฟล์ใน Microsoft Word หากต้องการเขียนทับชื่อไฟล์เดิม")

if __name__ == '__main__':
    base_dir = os.path.dirname(os.path.abspath(__file__))
    input_files = [
        os.path.join(base_dir, 'manuscript_part1_ch01_ch03.md'),
        os.path.join(base_dir, 'manuscript_part2_ch04_ch06.md'),
        os.path.join(base_dir, 'manuscript_part3_ch07_ch10.md'),
        os.path.join(base_dir, 'manuscript_part4_references_appendices.md')
    ]
    output_file = os.path.join(base_dir, 'Report_Final.docx')
    parse_manuscript_files(input_files, output_file)
