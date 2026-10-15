import docx

doc = docx.Document('Report_Corrected.docx')
for i, p in enumerate(doc.paragraphs):
    text = p.text.replace('\u200b', '')
    if 'จุดเริ่มต้น' in text or 'รูปที่ 10.2' in text or 'Concurrency Model' in text:
        print(f"P[{i}]: {len(text)} chars | {text[:80]}")

for i, t in enumerate(doc.tables):
    for r, row in enumerate(t.rows):
        for c, cell in enumerate(row.cells):
            text = cell.text.replace('\u200b', '')
            if 'จุดเริ่มต้น' in text or 'รูปที่ 10.2' in text or 'Concurrency Model' in text:
                print(f"Table[{i}] r={r} c={c}: {len(text)} chars | {text[:80]}")
